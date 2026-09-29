"""CSV/Excel connector: a price feed exported from anywhere.

The realistic onboarding path for a small retailer. Column names are matched
case-insensitively against a set of aliases so an analyst's spreadsheet works
without being reformatted first.
"""

from __future__ import annotations

import csv
import datetime as dt
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from app.core.logging import get_logger
from app.models.enums import StockState
from app.models.market import Competitor
from app.services.connectors.base import (
    ConnectorError,
    ConnectorResult,
    RawListing,
    RawObservation,
)

logger = get_logger(__name__)

#: Accepted spellings for each field. Lower-cased, non-alphanumerics stripped.
_ALIASES: dict[str, tuple[str, ...]] = {
    "external_id": ("externalid", "id", "sku", "productid", "itemid", "listingid"),
    "title": ("title", "name", "productname", "product", "description"),
    "brand": ("brand", "manufacturer", "make"),
    "gtin": ("gtin", "ean", "upc", "barcode", "isbn"),
    "url": ("url", "link", "producturl", "pageurl"),
    "price": ("price", "currentprice", "saleprice", "amount"),
    "shipping_cost": ("shipping", "shippingcost", "delivery", "deliverycost"),
    "currency": ("currency", "ccy"),
    "observed_at": ("observedat", "date", "timestamp", "capturedat", "scrapedat"),
    "stock_state": ("stock", "stockstate", "availability", "instock", "available"),
    "promo_flag": ("promo", "promoflag", "onsale", "discounted"),
}

_IN_STOCK_WORDS = {"in_stock", "instock", "available", "yes", "true", "1", "y"}
_OUT_OF_STOCK_WORDS = {"out_of_stock", "outofstock", "unavailable", "no", "false", "0", "n", "oos"}
_TRUE_WORDS = {"yes", "true", "1", "y", "promo", "sale"}


def _key(name: str) -> str:
    return "".join(ch for ch in name.lower() if ch.isalnum())


def _build_header_map(fieldnames: list[str]) -> dict[str, str]:
    """Map our field names to the spreadsheet's actual column names."""
    normalised = {_key(name): name for name in fieldnames}
    mapping: dict[str, str] = {}
    for field_name, aliases in _ALIASES.items():
        for alias in aliases:
            if alias in normalised:
                mapping[field_name] = normalised[alias]
                break
    return mapping


def _decimal(value: Any, default: str = "0") -> Decimal:
    text = str(value or "").strip().replace(",", "").replace("$", "")
    if not text:
        return Decimal(default)
    try:
        return Decimal(text)
    except InvalidOperation as exc:
        raise ConnectorError(f"Cannot parse {value!r} as a number.") from exc


def _stock(value: Any) -> StockState:
    text = _key(str(value or ""))
    if text in _IN_STOCK_WORDS:
        return StockState.IN_STOCK
    if text in _OUT_OF_STOCK_WORDS:
        return StockState.OUT_OF_STOCK
    return StockState.UNKNOWN


def _timestamp(value: Any, fallback: dt.datetime) -> dt.datetime:
    text = str(value or "").strip()
    if not text:
        return fallback
    try:
        parsed = dt.datetime.fromisoformat(text.replace("Z", "+00:00"))
    except ValueError:
        return fallback
    return parsed if parsed.tzinfo else parsed.replace(tzinfo=dt.UTC)


def _read_rows(path: Path) -> tuple[list[str], list[dict[str, Any]]]:
    if path.suffix.lower() in {".xlsx", ".xlsm"}:
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=True, data_only=True)
        sheet = workbook.active
        if sheet is None:
            raise ConnectorError(f"{path.name} has no readable worksheet.")
        rows = sheet.iter_rows(values_only=True)
        header = [str(cell) if cell is not None else "" for cell in next(rows, ())]
        records = [dict(zip(header, row, strict=False)) for row in rows]
        workbook.close()
        return header, records

    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        return list(reader.fieldnames or []), list(reader)


class CsvConnector:
    """Reads listings and prices from ``connector_config['path']``."""

    name = "csv"

    def __init__(self, base_dir: Path | None = None) -> None:
        #: Uploads are confined to this directory so a crafted config value
        #: cannot be used to read arbitrary files off the server.
        self.base_dir = base_dir

    def _resolve(self, raw_path: str) -> Path:
        path = Path(raw_path).expanduser()
        if self.base_dir is not None:
            base = self.base_dir.resolve()
            path = (base / path).resolve() if not path.is_absolute() else path.resolve()
            if not path.is_relative_to(base):
                raise ConnectorError(
                    f"Refusing to read {raw_path!r}: outside the upload directory."
                )
        else:
            path = path.resolve()
        if not path.is_file():
            raise ConnectorError(f"Feed file not found: {path}")
        return path

    def fetch(self, competitor: Competitor, *, now: dt.datetime) -> ConnectorResult:
        config = competitor.connector_config or {}
        raw_path = str(config.get("path", "")).strip()
        if not raw_path:
            raise ConnectorError(f"Competitor '{competitor.slug}' has no connector_config['path'].")

        path = self._resolve(raw_path)
        header, records = _read_rows(path)
        mapping = _build_header_map(header)

        missing = {"external_id", "price"} - mapping.keys()
        if missing:
            raise ConnectorError(
                f"{path.name} is missing required column(s): {', '.join(sorted(missing))}. "
                f"Found: {', '.join(header)}"
            )

        result = ConnectorResult()
        seen: set[str] = set()
        for line_no, record in enumerate(records, start=2):
            external_id = str(record.get(mapping["external_id"], "") or "").strip()
            if not external_id:
                result.errors.append(f"line {line_no}: blank identifier, skipped")
                continue
            try:
                price = _decimal(record.get(mapping["price"]))
            except ConnectorError as exc:
                result.errors.append(f"line {line_no}: {exc}")
                continue
            if price <= 0:
                result.errors.append(f"line {line_no}: non-positive price, skipped")
                continue

            def field(name: str, row: dict[str, Any] = record) -> Any:
                column = mapping.get(name)
                return row.get(column) if column else None

            if external_id not in seen:
                seen.add(external_id)
                result.listings.append(
                    RawListing(
                        external_id=external_id,
                        title=str(field("title") or external_id).strip(),
                        brand=str(field("brand") or "").strip(),
                        gtin=(str(field("gtin")).strip() or None) if field("gtin") else None,
                        url=str(field("url") or "").strip(),
                    )
                )

            result.observations.append(
                RawObservation(
                    external_id=external_id,
                    observed_at=_timestamp(field("observed_at"), now),
                    price=price,
                    shipping_cost=_decimal(field("shipping_cost")),
                    currency=str(field("currency") or "USD").strip()[:3].upper() or "USD",
                    stock_state=_stock(field("stock_state")),
                    promo_flag=_key(str(field("promo_flag") or "")) in _TRUE_WORDS,
                    raw_payload={"source": self.name, "file": path.name, "line": line_no},
                )
            )

        logger.info("CSV connector read %s: %s", path.name, result.summary())
        return result
